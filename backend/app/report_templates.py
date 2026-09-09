from __future__ import annotations

import hashlib
import re
import zipfile
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.shared import Inches

from .exporting import _chart_png


PLACEHOLDER = re.compile(r"\{\{\s*([a-zA-Z0-9_.-]{1,120})\s*\}\}")
MAX_TEMPLATE_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 2_000


class UnsafeTemplate(ValueError):
    pass


def inspect_docx_template(content: bytes) -> dict[str, Any]:
    """Validate an inert DOCX package and return deterministic template metadata."""
    if not content or len(content) > MAX_TEMPLATE_BYTES:
        raise UnsafeTemplate("模板不能为空且不得超过 10 MB")
    if not content.startswith(b"PK"):
        raise UnsafeTemplate("文件不是有效的 DOCX 压缩包")
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ARCHIVE_ENTRIES:
                raise UnsafeTemplate("模板包含过多内部文件")
            total = sum(item.file_size for item in infos)
            if total > MAX_UNCOMPRESSED_BYTES:
                raise UnsafeTemplate("模板解压后体积超过 100 MB")
            names = {item.filename.replace("\\", "/") for item in infos}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise UnsafeTemplate("DOCX 缺少必要的 Word 文档结构")
            for name in names:
                if name.startswith("/") or "../" in name.split("/"):
                    raise UnsafeTemplate("模板包含不安全的内部路径")
                lowered = name.casefold()
                if lowered.endswith("vbaproject.bin") or "/embeddings/" in lowered or lowered.endswith(".bin"):
                    raise UnsafeTemplate("模板包含宏或嵌入对象，已拒绝")
            for name in names:
                if not name.endswith(".rels"):
                    continue
                xml = archive.read(name).decode("utf-8", errors="ignore")
                if re.search(r'TargetMode\s*=\s*["\']External["\']', xml, re.I):
                    raise UnsafeTemplate("模板包含外部链接关系，已拒绝")
            text = "\n".join(
                archive.read(name).decode("utf-8", errors="ignore")
                for name in names
                if name.startswith("word/") and name.endswith(".xml")
            )
    except zipfile.BadZipFile as exc:
        raise UnsafeTemplate("DOCX 压缩结构损坏") from exc
    placeholders = sorted(set(PLACEHOLDER.findall(text)))
    return {
        "placeholders": placeholders,
        "placeholder_count": len(placeholders),
        "sha256": hashlib.sha256(content).hexdigest(),
        "size_bytes": len(content),
        "safe": True,
    }


def template_context(report: Any, custom_mapping: dict[str, Any] | None = None) -> dict[str, Any]:
    document = report.document or {}
    evidence = document.get("evidence") or []
    blocks = document.get("blocks") or []
    recommendations = document.get("metadata", {}).get("recommendations") or []
    context: dict[str, Any] = {
        "report.title": report.title,
        "report.summary": document.get("summary", ""),
        "report.generated_at": datetime.now(timezone.utc).astimezone().strftime("%Y-%m-%d %H:%M"),
        "report.body": "\n".join(str(block.get("content", "")) for block in blocks if block.get("kind") not in {"chart", "page_break"}),
        "recommendations": "\n".join(
            f"{index}. {item.get('recommendation', '')}" for index, item in enumerate(recommendations, 1)
        ),
        "limitations": "统计关系不等于因果关系；预测结果应结合业务变化持续回测。",
    }
    for item in evidence:
        evidence_id = str(item.get("id", ""))
        if evidence_id:
            context[f"evidence.{evidence_id}"] = item.get("value", "")
        if evidence_id.startswith("metric-"):
            context[f"metric.{evidence_id.removeprefix('metric-')}"] = item.get("value", "")
    for key, value in (custom_mapping or {}).items():
        if PLACEHOLDER.fullmatch("{{" + str(key) + "}}"):
            context[str(key)] = value
    return context


def render_docx_template(template_path: Path, report: Any, custom_mapping: dict[str, Any] | None = None) -> tuple[bytes, dict[str, Any]]:
    content = template_path.read_bytes()
    metadata = inspect_docx_template(content)
    word = Document(BytesIO(content))
    context = template_context(report, custom_mapping)
    charts = {str(item.get("id")): item for item in (report.document or {}).get("charts", [])}
    filled: set[str] = set()
    unresolved: set[str] = set()

    for paragraph in _all_paragraphs(word):
        keys = PLACEHOLDER.findall(paragraph.text)
        if not keys:
            continue
        exact = PLACEHOLDER.fullmatch(paragraph.text.strip())
        if exact and exact.group(1).startswith("chart."):
            key = exact.group(1)
            chart = charts.get(key.removeprefix("chart."))
            if chart:
                image = _chart_png(chart)
                if image:
                    _replace_paragraph_text(paragraph, "")
                    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
                    run.add_picture(image, width=Inches(6.2))
                    filled.add(key)
                    continue
            unresolved.add(key)
            continue
        rendered = paragraph.text
        for key in keys:
            token_pattern = re.compile(r"\{\{\s*" + re.escape(key) + r"\s*\}\}")
            if key in context and context[key] not in (None, ""):
                rendered = token_pattern.sub(str(context[key]), rendered)
                filled.add(key)
            else:
                unresolved.add(key)
        _replace_paragraph_text(paragraph, rendered)
    output = BytesIO()
    word.save(output)
    return output.getvalue(), {
        **metadata,
        "filled": sorted(filled),
        "unresolved": sorted(unresolved - filled),
        "context_keys": sorted(context),
    }


def _all_paragraphs(word: Document) -> Iterable[Any]:
    for paragraph in word.paragraphs:
        yield paragraph
    for table in word.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                for nested in cell.tables:
                    for nested_row in nested.rows:
                        for nested_cell in nested_row.cells:
                            yield from nested_cell.paragraphs
    for section in word.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            yield from part.paragraphs
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def _replace_paragraph_text(paragraph: Any, value: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)
