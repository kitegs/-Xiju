"""Generate and structurally verify a commercial DOCX from a CSV dataset."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import sys
from pathlib import Path

import pandas as pd
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.exporting import build_docx  # noqa: E402
from app.services import build_report, profile_dataframe  # noqa: E402


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("dataset", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--expected-sales", type=float)
    parser.add_argument("--expected-profit", type=float)
    parser.add_argument("--expected-rows", type=int)
    args = parser.parse_args()

    frame = pd.read_csv(args.dataset, low_memory=False)
    report = build_report(frame, profile_dataframe(frame), "Global Superstore 管理层经营分析", "生成证据化商业报告")
    evidence = {item.id: item for item in report.evidence}
    sales = float(evidence["metric-sales"].value.replace(",", ""))
    profit = float(evidence["metric-profit"].value.replace(",", ""))
    if args.expected_rows is not None:
        assert len(frame) == args.expected_rows, (len(frame), args.expected_rows)
    if args.expected_sales is not None:
        assert math.isclose(sales, args.expected_sales, abs_tol=.01), (sales, args.expected_sales)
    if args.expected_profit is not None:
        assert math.isclose(profit, args.expected_profit, abs_tol=.01), (profit, args.expected_profit)
    assert report.quality and report.quality.passed
    assert all(finding.evidence_ids for finding in report.findings)
    assert all(action.evidence_ids and action.validation_metric for action in report.actions)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_bytes(build_docx(report.title, report.model_dump(mode="json")))
    word = Document(args.output)
    assert word.paragraphs and word.tables
    assert len(word.inline_shapes) == len([chart for chart in report.charts if chart.chart_type != "kpi"])
    print(json.dumps({
        "output": str(args.output.resolve()), "dataset_sha256": hashlib.sha256(args.dataset.read_bytes()).hexdigest(),
        "rows": len(frame), "columns": len(frame.columns),
        "sales": sales, "profit": profit, "quality": report.quality.model_dump(mode="json"),
        "findings": len(report.findings), "actions": len(report.actions), "charts": len(report.charts),
        "docx_tables": len(word.tables), "docx_images": len(word.inline_shapes),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
