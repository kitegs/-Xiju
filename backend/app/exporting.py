from __future__ import annotations

import math
from io import BytesIO
from typing import Any

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
matplotlib.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"]
matplotlib.rcParams["axes.unicode_minus"] = False
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from matplotlib.ticker import FuncFormatter

from .schemas import ReportDocument
from .services import assess_report_quality


def build_docx(title: str, document: dict[str, Any]) -> bytes:
    """Render an export only after the evidence and presentation gate passes."""
    validated = ReportDocument.model_validate(document)
    quality = assess_report_quality(validated)
    if not quality.passed:
        failures = "；".join(item.title for item in quality.issues if item.severity == "error")
        raise ValueError(f"报告质量门禁未通过（{quality.score}/100）：{failures}")
    validated.quality = quality
    document = validated.model_dump(mode="json")
    output = BytesIO()
    word = Document()
    section = word.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(1); section.bottom_margin = Inches(1)
    section.left_margin = Inches(1); section.right_margin = Inches(1)
    section.header_distance = Inches(.492); section.footer_distance = Inches(.492)
    styles = word.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].paragraph_format.space_before = Pt(0)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.10
    _configure_heading(styles["Heading 1"], 16, "000000", 16, 8)
    _configure_heading(styles["Heading 2"], 13, "000000", 12, 6)
    _configure_heading(styles["Heading 3"], 12, "000000", 8, 4)
    _configure_heading(styles["Title"], 24, "000000", 0, 8)
    _remove_style_borders(styles["Title"])

    kicker = word.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(8)
    kicker_run = kicker.add_run("析据 Xiju  数据分析报告")
    _set_run(kicker_run, 10, "505A66", True)
    heading = word.add_paragraph(style="Title")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = heading.add_run(title)
    _set_run(title_run, 24, "000000", True)
    summary = word.add_paragraph(document.get("summary") or "")
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.space_after = Pt(8)
    for run in summary.runs: _set_run(run, 11, "505A66", False)
    metadata = document.get("metadata") or {}
    scope = word.add_paragraph(f"数据范围  {metadata.get('row_count', 0):,} 行  {metadata.get('column_count', 0):,} 个字段    自动规则检查  {quality.score}/100（不代表业务结论已获确认）")
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope.paragraph_format.space_after = Pt(16)
    for run in scope.runs: _set_run(run, 9, "667085", False)
    charts = {item.get("id"): item for item in document.get("charts", [])}
    kpi_charts = [item for item in charts.values() if item.get("chart_type") == "kpi"]
    rendered_kpis = False

    for block in document.get("blocks", []):
        kind = block.get("kind")
        if kind == "heading":
            paragraph = word.add_heading(block.get("content") or "", level=1)
            paragraph.paragraph_format.keep_with_next = True
        elif kind == "page_break":
            word.add_page_break()
        elif kind == "chart":
            chart = charts.get(block.get("chart_id"))
            if chart:
                if chart.get("chart_type") == "kpi":
                    if not rendered_kpis:
                        _add_kpi_table(word, kpi_charts)
                        rendered_kpis = True
                    continue
                chart_heading = word.add_heading(chart.get("title") or "图表", level=2)
                chart_heading.paragraph_format.keep_with_next = True
                image = _chart_png(chart)
                if image:
                    word.add_picture(image, width=Inches(6.05))
                    picture = word.paragraphs[-1]
                    picture.paragraph_format.keep_together = True
                    picture.paragraph_format.keep_with_next = True
                    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    doc_properties = word.inline_shapes[-1]._inline.docPr
                    doc_properties.set("title", chart.get("title") or "数据图表")
                    doc_properties.set("descr", chart.get("description") or f"{chart.get('title', '数据图表')}，数据与证据详见后续附录。")
                else:
                    _add_data_table(word, chart)
                if chart.get("description"):
                    paragraph = word.add_paragraph(chart["description"])
                    paragraph.style = word.styles["Caption"]
                    paragraph.paragraph_format.keep_together = True
                    paragraph.paragraph_format.keep_with_next = True
                _add_source_note(word, chart)
        elif block.get("id") == "actions" and document.get("actions"):
            _add_action_table(word, document["actions"])
        else:
            headline = (block.get("style") or {}).get("headline")
            if headline:
                paragraph = word.add_paragraph()
                paragraph.paragraph_format.keep_together = True
                run = paragraph.add_run(headline + "。")
                _set_run(run, 11, "000000", True)
                body = block.get("content") or ""
                if body.startswith(headline):
                    body = body[len(headline):].lstrip("。.!！?？ \n")
                if body:
                    body_run = paragraph.add_run(" " + body.replace("\n", " "))
                    _set_run(body_run, 11, "202124", False)
            else:
                paragraph = word.add_paragraph(block.get("content") or "")
                paragraph.paragraph_format.keep_together = kind in {"insight", "callout"}

    evidence = document.get("evidence") or []
    requirements = metadata.get('report_requirements') or {}
    if requirements:
        visible = {key for block in document.get('blocks', []) for key in block.get('evidence_ids', [])}
        evidence = [item for item in evidence if item.get('id') in visible]
    if evidence and metadata.get('appendix_visible', True):
        word.add_heading("证据附录", level=1).paragraph_format.page_break_before = True
        word.add_paragraph(f"本报告通过自动质量检查，得分 {quality.score}/100。下表列出正文数字和图表使用的本地计算证据。")
        _add_evidence_table(word, evidence)
        coded = [item for item in evidence if item.get("code")]
        if coded:
            word.add_heading("复现代码", level=2).paragraph_format.page_break_before = True
            for item in coded:
                label = word.add_paragraph()
                label.paragraph_format.keep_with_next = True
                run = label.add_run(f"{item.get('id')}  {item.get('statement')}")
                _set_run(run, 9, "000000", True)
                code = word.add_paragraph(item["code"])
                code.style = word.styles["No Spacing"]
                code.paragraph_format.space_after = Pt(3)
                for run in code.runs: _set_run(run, 8, "202124", False)
                _shade(code, "F4F5F7")

    footer = section.footer.paragraphs[0]
    footer.text = "析据 Xiju  自动生成报告  发布前需复核    第 "
    _add_field(footer, "PAGE")
    footer.add_run(" 页")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    word.save(output)
    return output.getvalue()


def _chart_png(chart: dict[str, Any]) -> BytesIO | None:
    chart_type = chart.get("chart_type")
    rows = chart.get("data") or []
    if not rows or chart_type in {"table", "highlight_table", "heatmap", "boxplot"}:
        return None
    x_key = (chart.get("x") or {}).get("column") or "label"
    series = [item.get("column") for item in chart.get("series", []) if item.get("column")]
    y_key = (chart.get("y") or {}).get("column") or "value"
    series = list(dict.fromkeys(series or [y_key]))
    labels = [str(row.get(x_key, row.get("label", ""))) for row in rows]
    fig, axis = plt.subplots(figsize=(7.2, 3.25), dpi=150)
    colors = ["#3B6FB6", "#E28E2C", "#5A9A72", "#9B6FB6"]
    try:
        if chart_type == "kpi":
            axis.axis("off")
            value = rows[0].get(y_key, rows[0].get("value"))
            suffix = rows[0].get("suffix", "")
            axis.text(.5, .58, f"{float(value):,.2f}{suffix}" if value is not None else "—", ha="center", va="center", fontsize=24, color="#354052")
            axis.text(.5, .38, labels[0], ha="center", va="center", fontsize=12, color="#7A8798")
        elif chart_type in {"scatter", "bubble"}:
            x_values = [float(row.get(x_key, 0) or 0) for row in rows]
            y_values = [float(row.get(y_key, 0) or 0) for row in rows]
            axis.scatter(x_values, y_values, s=12 if chart_type == "scatter" else 24, alpha=.45, color=colors[0])
            axis.set_xlabel(x_key); axis.set_ylabel(y_key)
        elif chart_type in {"pie", "donut"}:
            values = [float(row.get(y_key, row.get("value", 0)) or 0) for row in rows[:10]]
            axis.pie(values, labels=labels[:10], autopct="%1.1f%%", startangle=90, wedgeprops={"width": .48 if chart_type == "donut" else 1})
        elif chart_type == "bar" and (chart.get("style") or {}).get("orientation") == "horizontal":
            values = [float(row.get(series[0], 0) or 0) for row in rows]
            axis.barh(labels, values, color=colors[0])
            axis.invert_yaxis()
            if (chart.get("style") or {}).get("show_labels"):
                for label, value in zip(labels, values): axis.text(value, label, " " + _compact_number(value), va="center", fontsize=8)
            axis.xaxis.set_major_formatter(FuncFormatter(lambda value, _: _compact_number(value)))
        else:
            positions = list(range(len(labels)))
            width = .8 / max(len(series), 1)
            for index, key in enumerate(series):
                values = [float(row.get(key, 0) or 0) for row in rows]
                if chart_type in {"line", "area", "combo"}:
                    axis.plot(positions, values, marker="o", linewidth=2, label=key, color=colors[index % len(colors)])
                    if chart_type == "area": axis.fill_between(positions, values, alpha=.15, color=colors[index % len(colors)])
                else:
                    axis.bar([position + (index - (len(series)-1)/2) * width for position in positions], values, width=width, label=key, color=colors[index % len(colors)])
            step = max(1, math.ceil(len(labels) / 10))
            shown = [label if index % step == 0 or index == len(labels) - 1 else "" for index, label in enumerate(labels)]
            axis.set_xticks(positions, shown, rotation=30, ha="right")
            if len(series) > 1: axis.legend(frameon=False)
            axis.yaxis.set_major_formatter(FuncFormatter(lambda value, _: _compact_number(value)))
        axis.grid(axis="y", alpha=.18)
        fig.tight_layout()
        image = BytesIO(); fig.savefig(image, format="png", bbox_inches="tight", facecolor="white"); image.seek(0)
        return image
    finally:
        plt.close(fig)


def _add_data_table(word: Document, chart: dict[str, Any]) -> None:
    rows = chart.get("data") or []
    if not rows:
        word.add_paragraph("暂无可导出的图表数据。")
        return
    columns = list(rows[0].keys())[:8]
    table = word.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    widths = _column_widths(columns)
    _set_table_geometry(table, widths)
    for index, column in enumerate(columns):
        table.rows[0].cells[index].text = str(column)
        _shade_cell(table.rows[0].cells[index], "F2F4F7")
        for run in table.rows[0].cells[index].paragraphs[0].runs: run.bold = True
    for item in rows[:30]:
        cells = table.add_row().cells
        for index, column in enumerate(columns):
            cells[index].text = str(item.get(column, ""))
            cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _set_table_geometry(table, widths)
    _set_table_borders(table)


def _add_kpi_table(word: Document, charts: list[dict[str, Any]]) -> None:
    if not charts:
        return
    table = word.add_table(rows=1, cols=len(charts))
    table.autofit = False
    for index, chart in enumerate(charts):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        value = (chart.get("data") or [{}])[0]
        amount = value.get("value")
        suffix = value.get("suffix", "")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label = paragraph.add_run((chart.get("title") or "指标") + "\n")
        _set_run(label, 9, "667085", False)
        displayed = f"{float(amount):,.2f}" if amount is not None and suffix == "%" else (_compact_number(float(amount)) if amount is not None else "—")
        number = paragraph.add_run(displayed + suffix)
        _set_run(number, 17, "203748", True)
        _shade_cell(cell, "F6F8FB")
    _set_table_geometry(table, [round(9360 / len(charts))] * len(charts))
    _set_table_borders(table)
    word.add_paragraph().paragraph_format.space_after = Pt(2)


def _add_source_note(word: Document, chart: dict[str, Any]) -> None:
    style = chart.get("style") or {}
    parts = []
    if style.get("source_label"): parts.append("来源：" + style["source_label"])
    if style.get("unit"): parts.append("单位：" + style["unit"])
    if chart.get("evidence_ids"): parts.append("证据：" + "、".join(chart["evidence_ids"]))
    if parts:
        paragraph = word.add_paragraph("  ".join(parts))
        paragraph.paragraph_format.space_after = Pt(8)
        for run in paragraph.runs: _set_run(run, 8, "667085", False)


def _add_action_table(word: Document, actions: list[dict[str, Any]]) -> None:
    table = word.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["优先级", "行动", "依据", "验证指标"]
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]; cell.text = header; _shade_cell(cell, "203748")
        for run in cell.paragraphs[0].runs: _set_run(run, 9, "FFFFFF", True)
    labels = {"high": "高", "medium": "中", "low": "低"}
    for row_index, item in enumerate(actions):
        cells = table.add_row().cells
        values = [labels.get(item.get("priority"), "中"), item.get("action", ""), item.get("rationale", ""), item.get("validation_metric", "")]
        for index, value in enumerate(values):
            cells[index].text = str(value); cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2: _shade_cell(cells[index], "F6F8FB")
            for run in cells[index].paragraphs[0].runs: _set_run(run, 9, "202124", False)
        _prevent_row_split(table.rows[-1])
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    _set_table_geometry(table, [900, 3000, 3000, 2460]); _set_table_borders(table)


def _add_evidence_table(word: Document, evidence: list[dict[str, Any]]) -> None:
    # Long lineage lists belong in full-width prose, not in a 1-inch cell that
    # forces a nearly page-high unbreakable row and a mostly blank previous page.
    lengthy = [item for item in evidence if len("、".join(item.get("source_columns") or [])) > 140]
    for item in lengthy:
        heading = word.add_paragraph()
        heading.paragraph_format.keep_with_next = True
        _set_run(heading.add_run(f"{item.get('id', '')}  {item.get('statement', '')}"), 10, "000000", True)
        for text in (item.get("value", ""), "方法：" + item.get("method", ""),
                     "来源字段：" + "、".join(item.get("source_columns") or [])):
            paragraph = word.add_paragraph(text)
            for run in paragraph.runs:
                _set_run(run, 10, "202124", False)
    evidence = [item for item in evidence if item not in lengthy]
    if not evidence:
        return
    table = word.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["编号", "结论与结果", "方法", "来源字段"]
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]; cell.text = header; _shade_cell(cell, "203748")
        for run in cell.paragraphs[0].runs: _set_run(run, 8.5, "FFFFFF", True)
    for row_index, item in enumerate(evidence):
        cells = table.add_row().cells
        values = [item.get("id", ""), f"{item.get('statement', '')}\n{item.get('value', '')}", item.get("method", ""), "、".join(item.get("source_columns") or []) or "完整数据集"]
        for index, value in enumerate(values):
            cells[index].text = str(value); cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2: _shade_cell(cells[index], "F6F8FB")
            for run in cells[index].paragraphs[0].runs: _set_run(run, 8.5, "202124", False)
        _prevent_row_split(table.rows[-1])
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    _set_table_geometry(table, [1500, 3300, 3000, 1560]); _set_table_borders(table)


def _shade(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), fill); properties.append(shading)


def _configure_heading(style, size: float, color: str, before: float, after: float) -> None:
    style.font.name = "Microsoft YaHei"; style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string(color)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)


def _remove_style_borders(style) -> None:
    properties = style._element.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is not None:
        properties.remove(borders)


def _set_run(run, size: float, color: str, bold: bool) -> None:
    run.font.name = "Microsoft YaHei"; run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color); run.bold = bold
    run._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def _column_widths(columns: list[str]) -> list[int]:
    weights = [max(1.0, min(3.0, len(str(column)) / 6)) for column in columns]
    total = sum(weights)
    widths = [round(9360 * weight / total) for weight in weights]
    widths[-1] += 9360 - sum(widths)
    return widths


def _set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    tbl = table._tbl; properties = tbl.tblPr
    width = properties.first_child_found_in("w:tblW")
    if width is None: width = OxmlElement("w:tblW"); properties.append(width)
    width.set(qn("w:type"), "dxa"); width.set(qn("w:w"), "9360")
    indent = properties.first_child_found_in("w:tblInd")
    if indent is None: indent = OxmlElement("w:tblInd"); properties.append(indent)
    indent.set(qn("w:type"), "dxa"); indent.set(qn("w:w"), "120")
    grid = tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol"); column.set(qn("w:w"), str(value)); grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_width is None: tc_width = OxmlElement("w:tcW"); cell._tc.get_or_add_tcPr().append(tc_width)
            tc_width.set(qn("w:type"), "dxa"); tc_width.set(qn("w:w"), str(widths[index]))
            margins = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcMar")
            if margins is None: margins = OxmlElement("w:tcMar"); cell._tc.get_or_add_tcPr().append(margins)
            for side, value in (("top",80),("bottom",80),("start",120),("end",120)):
                node = margins.find(qn(f"w:{side}"))
                if node is None: node = OxmlElement(f"w:{side}"); margins.append(node)
                node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def _set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None: borders = OxmlElement("w:tblBorders"); properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None: node = OxmlElement(f"w:{edge}"); borders.append(node)
        node.set(qn("w:val"), "single"); node.set(qn("w:sz"), "4"); node.set(qn("w:color"), "D9D9D9")


def _prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def _add_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText"); instruction.set(qn("xml:space"), "preserve"); instruction.text = field
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def _compact_number(value: float) -> str:
    absolute = abs(value)
    if absolute >= 100_000_000: return f"{value / 100_000_000:.1f}亿"
    if absolute >= 10_000: return f"{value / 10_000:.1f}万"
    if absolute >= 1_000: return f"{value / 1_000:.1f}千"
    if absolute >= 10: return f"{value:,.0f}"
    return f"{value:,.2f}"


def _shade_cell(cell, fill: str) -> None:
    shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), fill); cell._tc.get_or_add_tcPr().append(shading)
