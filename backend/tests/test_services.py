import pandas as pd
from docx import Document
from sqlalchemy import create_engine, inspect

from app.schemas import CleaningStep
from app.external_bi import chart_to_superset_params
from app.services import apply_cleaning_steps, assess_report_quality, build_report, profile_dataframe, quality_summary, run_sql_query
from app.exporting import build_docx
from app.forecasting import forecast_time_series, recommendation_actions
from app.report_templates import UnsafeTemplate, inspect_docx_template, render_docx_template
from app.research import inferential_statistics, survey_cross_analysis
from app.policy import ToolPolicyError, decide_tool, enforce_tools
from app.prompting import planning_context, validate_synthesis
from app.main import _validate_model_plan
from app.database import ensure_compat_schema


def test_usage_log_keeps_stage_links_and_marks_missing_provider_usage():
    from app.llm import CompletionResult
    from app.main import usage_log
    from app.models import ProviderConfig

    provider = ProviderConfig(
        workspace_id="workspace", provider="deepseek", enabled=True, is_default=True,
        base_url="https://api.deepseek.com", model="deepseek-test",
    )
    result = CompletionResult(
        content="{}", request_id="request", model="deepseek-test",
        usage={"prompt_tokens": 0, "cache_hit_tokens": 0, "cache_miss_tokens": 0, "completion_tokens": 0, "total_tokens": 0},
        estimated_cost_cny=0, latency_ms=12, pricing={}, usage_available=False,
    )
    row = usage_log(
        "workspace", "conversation", "plan", provider, result,
        stage="synthesis", purpose="整理运行证据结论", run_id="run", report_id="report",
    )
    assert row.stage == "synthesis"
    assert row.run_id == "run" and row.report_id == "report"
    assert row.usage_unavailable is True
    assert row.total_tokens == 0


def test_analysis_intake_asks_once_and_tracks_recommended_defaults():
    from app.intake import apply_brief_update, evaluate_intake
    from app.schemas import AnalysisBrief, AnalysisBriefUpdate

    profile = {"columns": [{"name": "Year"}, {"name": "Sales"}, {"name": "Profit"}]}
    decision = evaluate_intake("生成经营报告", profile, AnalysisBrief(), "auto")
    assert decision.should_pause is True
    assert len(decision.questions) == 8
    assert len({item.question_key for item in decision.questions}) == 8

    brief = apply_brief_update(AnalysisBrief(), AnalysisBriefUpdate(
        workspace_id="workspace", objective="生成经营报告",
        answers={"audience": "总经理"}, use_recommended_defaults=True,
    ), decision)
    assert brief.answers["audience"] == "总经理"
    assert "audience" in brief.confirmed_keys
    assert "audience" not in brief.defaulted_keys
    assert len(brief.defaulted_keys) == 7

    repeated = evaluate_intake("生成经营报告", profile, brief, "always")
    assert repeated.should_pause is False
    assert repeated.questions == []
    assert repeated.status == "completed"

    disabled = evaluate_intake("生成另一份报告", profile, AnalysisBrief(), "off")
    assert disabled.should_pause is False
    assert disabled.status == "disabled"


def test_build_report_keeps_evidence_and_editable_chart_spec():
    frame = pd.DataFrame({"区域": ["华东", "华东", "华南"], "销售额": [10, 20, 15]})
    profile = profile_dataframe(frame)

    document = build_report(frame, profile, "销售分析", "找出主要来源")

    assert document.title == "销售分析"
    assert document.evidence
    assert document.charts[0].chart_type == "bar"
    assert document.charts[0].x.column == "区域"
    assert document.charts[0].y.aggregate == "sum"
    assert document.charts[0].evidence_ids


def test_profile_reports_shape_and_numeric_summary():
    profile = profile_dataframe(pd.DataFrame({"value": [1, 2, None], "label": ["a", "b", "c"]}))

    assert profile["row_count"] == 3
    assert profile["column_count"] == 2
    assert profile["columns"][0]["null_count"] == 1
    assert profile["columns"][0]["mean"] == 1.5


def test_quality_and_cleaning_steps_are_deterministic():
    frame = pd.DataFrame({"地区": [" 华东 ", " 华东 ", None], "销售额": [10, 10, None]})
    quality = quality_summary(frame)
    assert quality.missing_cells == 2
    assert quality.duplicate_rows == 1

    cleaned, step_results, changed_cells = apply_cleaning_steps(frame, [
        CleaningStep(operation="trim_text", column="地区"),
        CleaningStep(operation="fill_missing", column="地区", method="constant", value="未知"),
        CleaningStep(operation="fill_missing", column="销售额", method="mean"),
        CleaningStep(operation="drop_duplicates"),
    ])
    assert len(cleaned) == 2
    assert cleaned.iloc[0]["地区"] == "华东"
    assert cleaned.iloc[1]["销售额"] == 10
    assert step_results[-1]["removed_rows"] == 1
    assert changed_cells >= 3


def test_global_superstore_style_report_uses_business_fields_and_full_evidence():
    frame = pd.DataFrame({
        "Region": ["West", "East", "West"], "Category": ["Office", "Tech", "Tech"],
        "Sales": [100.0, 200.0, 300.0], "Profit": [10.0, -20.0, 60.0],
        "Discount": [0.0, 0.5, 0.1], "Quantity": [1, 2, 3],
        "Order Date": ["00:00.0"] * 3,
    })
    profile = profile_dataframe(frame)
    quality = quality_summary(frame)
    document = build_report(frame, profile, "经营报告", "生成管理层报告")

    assert len(profile["columns"]) == len(frame.columns)
    assert profile["read_mode"] == "full" and profile["is_truncated"] is False
    assert {issue.issue_type for issue in quality.issues} >= {"constant", "invalid_date"}
    assert document.charts[0].x.column == "Region"
    assert [item.column for item in document.charts[0].series] == ["Sales"]
    assert next(item for item in document.evidence if item.id == "metric-sales").value == "600.00"
    assert next(item for item in document.evidence if item.id == "metric-margin").value == "8.33%"
    assert document.findings and document.actions
    assert document.quality and document.quality.passed


def test_report_quality_gate_rejects_ungrounded_numbers_and_category_lines():
    frame = pd.DataFrame({"Region": ["West", "East"], "Sales": [100, 200]})
    document = build_report(frame, profile_dataframe(frame), "经营报告", "生成报告")
    assert assess_report_quality(document).passed

    document.charts[0].chart_type = "line"
    document.blocks.append(type(document.blocks[0])(
        id="manual-number", kind="paragraph", content="销售额增长了 25%。",
    ))
    checked = assess_report_quality(document)
    issue_ids = {item.id for item in checked.issues}
    assert checked.passed is False
    assert "chart-primary-comparison-trend-non-temporal-dimension" in issue_ids
    assert "ungrounded-number-manual-number" in issue_ids


def test_business_docx_is_compact_and_contains_quality_and_action_tables(tmp_path):
    frame = pd.DataFrame({
        "Region": ["West", "East", "West"], "Sales": [100, 200, 300],
        "Profit": [10, -20, 60], "Discount": [0.0, 0.5, 0.1], "Quantity": [1, 2, 3],
    })
    document = build_report(frame, profile_dataframe(frame), "经营报告", "生成管理层报告")
    content = build_docx(document.title, document.model_dump(mode="json"))
    output = tmp_path / "report.docx"; output.write_bytes(content)
    word = Document(output)
    text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    assert "质量检查" in text and "建议行动" in text and "证据附录" in text
    group_title = next(chart.title for chart in document.charts if chart.id == "primary-comparison")
    assert f"{group_title}。 {group_title}" not in text
    assert len(word.inline_shapes) == len([chart for chart in document.charts if chart.chart_type != "kpi"])
    assert len(word.tables) >= 3


def test_duckdb_copy_query_is_bounded_and_blocks_external_or_writing_sql():
    frame = pd.DataFrame({"Region": ["West", "East", "West"], "Sales": [10, 20, 30]})
    columns, rows, count, truncated, duration = run_sql_query(
        frame, 'SELECT "Region", SUM("Sales") AS total FROM dataset GROUP BY 1 ORDER BY total DESC', 10,
    )
    assert columns == ["Region", "total"]
    assert rows[0] == {"Region": "West", "total": 40.0}
    assert count == 2 and truncated is False and duration >= 1
    for unsafe in ("CREATE TABLE copy AS SELECT * FROM dataset", "ATTACH 'other.db' AS x", "COPY dataset TO 'x.csv'"):
        try:
            run_sql_query(frame, unsafe, 10)
        except ValueError as exc:
            assert "只允许 SELECT/WITH" in str(exc)
        else:
            raise AssertionError(f"unsafe SQL was accepted: {unsafe}")


def test_model_sql_quotes_known_columns_with_spaces():
    from app.services import normalize_sql_columns
    frame = pd.DataFrame({"Order Date": pd.to_datetime(["2024-01-01"]), "Sales": [10]})
    sql = normalize_sql_columns("SELECT MONTH(Order Date) AS month, SUM(Sales) FROM dataset GROUP BY 1", frame.columns)
    assert 'MONTH("Order Date")' in sql
    columns, rows, *_ = run_sql_query(frame, sql)
    assert columns == ["month", "sum(Sales)"] and rows[0]["month"] == 1


def test_convert_date_accepts_explicit_format_without_touching_source():
    from app.schemas import CleaningStep
    from app.services import apply_cleaning_steps

    source = pd.DataFrame({"when": ["01/02/2011", "02/03/2011"]})
    original = source.copy(deep=True)
    result, _, _ = apply_cleaning_steps(source, [CleaningStep(
        operation="convert_type", column="when", target_type="date", date_format="%m/%d/%Y",
    )])
    assert result["when"].tolist() == ["2011-01-02", "2011-02-03"]
    assert source.equals(original)


def test_superset_mapping_uses_mysql_metrics_and_supported_scatter():
    viz_type, params = chart_to_superset_params({
        "id": "sales", "title": "总销售额", "chart_type": "kpi",
        "y": {"column": "Sales", "aggregate": "sum"}, "series": [{"column": "Sales", "aggregate": "sum"}],
    }, 11)
    assert viz_type == "big_number_total"
    assert params["metric"]["sqlExpression"] == "SUM(`Sales`)"

    viz_type, params = chart_to_superset_params({
        "id": "scatter", "title": "折扣与利润", "chart_type": "scatter",
        "x": {"column": "Discount"}, "y": {"column": "Profit", "aggregate": None},
    }, 11)
    assert viz_type == "echarts_timeseries_scatter"
    assert params["x_axis"] == "Discount"
    assert params["metrics"][0]["sqlExpression"] == "AVG(`Profit`)"


def test_forecast_backtests_and_returns_prediction_intervals():
    dates = pd.date_range("2024-01-01", periods=18, freq="MS")
    frame = pd.DataFrame({"Order Date": dates, "Sales": [100 + 5 * index for index in range(18)]})
    result = forecast_time_series(frame, date_column="Order Date", target_column="Sales", horizon=4)
    assert result.method in {"naive", "moving_average", "linear_trend", "seasonal_naive"}
    assert len(result.forecast) == 4
    assert result.forecast[0]["lower"] <= result.forecast[0]["forecast"] <= result.forecast[0]["upper"]
    assert result.metrics["mae"] is not None
    assert result.diagnostics["backtest_folds"] >= 1
    assert "trend_per_period" in result.diagnostics
    assert "deterministic forecast" in result.code
    assert recommendation_actions(frame, quality_summary(frame), result)[0]["validation"]


def test_research_and_survey_tools_keep_effect_sizes_and_reproducible_methods():
    frame = pd.DataFrame({
        "组别": ["A"] * 8 + ["B"] * 8 + ["C"] * 8,
        "性别": ["F", "M"] * 12,
        "前测": list(range(1, 25)),
        "后测": [value * 1.5 for value in range(1, 25)],
        "Q1": [1, 2, 3, 4] * 6,
    })
    evidence = inferential_statistics(frame)
    ids = {item["id"] for item in evidence}
    assert {"research-correlation", "research-anova", "research-chi-square", "research-sample-size"} <= ids
    assert any("η²" in item["value"] for item in evidence)
    assert all(item["code"] for item in evidence)
    survey = survey_cross_analysis(frame)
    assert survey and "Cramér's V" in survey[0]["value"]


def test_docx_template_is_checked_and_filled_without_changing_original(tmp_path):
    template_path = tmp_path / "template.docx"
    word = Document(); word.add_heading("{{ report.title }}", level=1); word.add_paragraph("{{ report.summary }}"); word.add_paragraph("{{ metric.sales }}")
    word.save(template_path)
    inspected = inspect_docx_template(template_path.read_bytes())
    assert {"report.title", "report.summary", "metric.sales"} <= set(inspected["placeholders"])
    frame = pd.DataFrame({"Region": ["West", "East"], "Sales": [100, 200]})
    document = build_report(frame, profile_dataframe(frame), "模板报告", "测试")
    report = type("ReportStub", (), {"title": "模板报告", "document": document.model_dump(mode="json")})()
    rendered, metadata = render_docx_template(template_path, report)
    output_path = tmp_path / "rendered.docx"; output_path.write_bytes(rendered)
    text = "\n".join(paragraph.text for paragraph in Document(output_path).paragraphs)
    assert "模板报告" in text and "{{ report.title }}" not in text
    assert "report.title" in metadata["filled"]
    assert "{{ report.title }}" in "\n".join(paragraph.text for paragraph in Document(template_path).paragraphs)
    try:
        inspect_docx_template(b"not-a-docx")
    except UnsafeTemplate:
        pass
    else:
        raise AssertionError("unsafe file was accepted")


def test_execution_permission_matrix_is_enforced_server_side():
    expected = {
        "safe": {
            "dataset.profile": "allow", "sql.query": "approval", "python.run": "forbid",
            "data.clean.apply": "forbid", "mcp.call": "forbid", "superset.publish": "forbid",
        },
        "partial": {
            "dataset.profile": "allow", "sql.query": "allow", "python.run": "approval",
            "data.clean.apply": "approval", "mcp.call": "approval", "superset.publish": "approval",
        },
        "full": {
            "dataset.profile": "allow", "sql.query": "allow", "python.run": "allow",
            "data.clean.apply": "allow", "mcp.call": "allow", "superset.publish": "approval",
        },
    }
    for mode, tools in expected.items():
        assert {tool: decide_tool(mode, tool).decision for tool in tools} == tools
    for mode in expected:
        assert decide_tool(mode, "filesystem.delete").decision == "forbid"
    for mode, tool in (("safe", "python.run"), ("partial", "python.run"), ("full", "superset.publish")):
        try:
            enforce_tools(mode, [tool], approved=False)
        except ToolPolicyError:
            pass
        else:
            raise AssertionError(f"unauthorized tool call was accepted: {mode}/{tool}")
    assert enforce_tools("partial", ["python.run"], approved=True)[0].decision == "approval"


def test_context_packet_and_claim_validation_prevent_context_and_evidence_leaks():
    profile = {"row_count": 2, "column_count": 1, "columns": [{"name": "Sales", "dtype": "int64", "sample_values": [10, 20]}]}
    packet = planning_context(dataset_name="untrusted name", dataset_version_id="version-1", profile=profile)
    assert "CONTEXT_PACKET_V1" in packet and "trusted_state" in packet and "untrusted_dataset_metadata" in packet
    assert "full_dataframe" not in packet
    valid = validate_synthesis({
        "summary": "已完成证据化分析。",
        "claims": [{"text": "销售额为 {{total}}。", "type": "fact", "evidence_ids": ["sales-total"]}],
        "bindings": [{"name": "total", "evidence_id": "sales-total", "path": "value"}],
        "limitations": [],
    }, [{"id": "sales-total", "value": "30"}])
    assert valid.claims[0].evidence_ids == ["sales-total"]
    for raw in (
        {"summary": "销售额为 30。", "claims": [], "limitations": []},
        {"summary": "已完成。", "claims": [{"text": "销售额为 30。", "type": "fact", "evidence_ids": []}], "limitations": []},
        {"summary": "已完成。", "claims": [{"text": "销售额已确认。", "type": "fact", "evidence_ids": ["missing"]}], "limitations": []},
    ):
        try:
            validate_synthesis(raw, [{"id": "sales-total", "value": "30"}])
        except ValueError:
            pass
        else:
            raise AssertionError("ungrounded claim was accepted")


def test_model_plan_drops_unknown_tools_and_nonexistent_fields():
    profile = {"columns": [{"name": "Region"}, {"name": "Sales"}]}
    plan = _validate_model_plan({"steps": [
        {"tool": "filesystem.delete", "arguments": {}},
        {"tool": "timeseries.forecast", "arguments": {"date_column": "Imaginary Date", "target_column": "Sales"}},
        {"tool": "data.clean.apply", "arguments": {"steps": [{"operation": "trim_text", "column": "Missing"}]}},
        {"tool": "data.clean.apply", "arguments": {"steps": [{"operation": "trim_text", "column": "Region"}]}},
        {"tool": "sql.query", "arguments": {"sql": 'SELECT "Missing" FROM dataset'}},
        {"tool": "sql.query", "arguments": {"sql": 'SELECT "Region", SUM("Sales") AS "total" FROM dataset GROUP BY 1'}},
    ]}, "分析销售", "dataset-id", profile=profile)
    assert [step.tool for step in plan.steps].count("sql.query") == 1
    assert [step.tool for step in plan.steps].count("data.clean.apply") == 1
    assert plan.steps[-1].tool == "assistant.synthesize"
    assert all(step.tool != "filesystem.delete" for step in plan.steps)
    assert "Missing" not in str(plan.model_dump()) and "Imaginary Date" not in str(plan.model_dump())


def test_legacy_sqlite_schema_gets_durable_run_columns(tmp_path):
    legacy = create_engine(f"sqlite:///{tmp_path / 'legacy.db'}")
    with legacy.begin() as connection:
        connection.exec_driver_sql("CREATE TABLE projects (id VARCHAR(36) PRIMARY KEY)")
        connection.exec_driver_sql("CREATE TABLE reports (id VARCHAR(36) PRIMARY KEY)")
        connection.exec_driver_sql("CREATE TABLE conversations (id VARCHAR(36) PRIMARY KEY)")
        connection.exec_driver_sql("""
            CREATE TABLE llm_usage_logs (
                id VARCHAR(36) PRIMARY KEY, workspace_id VARCHAR(36), provider VARCHAR(40),
                model VARCHAR(160)
            )
        """)
        connection.exec_driver_sql("""
            CREATE TABLE analysis_runs (
                id VARCHAR(36) PRIMARY KEY, workspace_id VARCHAR(36), conversation_id VARCHAR(36),
                plan_message_id VARCHAR(36), status VARCHAR(24)
            )
        """)
        connection.exec_driver_sql("INSERT INTO analysis_runs (id,status) VALUES ('retained','interrupted')")
        ensure_compat_schema(connection)
        ensure_compat_schema(connection)
        assert connection.exec_driver_sql("SELECT status FROM analysis_runs WHERE id='retained'").scalar() == 'interrupted'
    columns = {item["name"] for item in inspect(legacy).get_columns("analysis_runs")}
    assert {"project_id", "dataset_version_ids", "analysis_spec", "idempotency_key", "approval_granted"} <= columns
    indexes = {item["name"] for item in inspect(legacy).get_indexes("analysis_runs")}
    assert "ix_analysis_runs_idempotency_key" in indexes
    project_columns = {item["name"] for item in inspect(legacy).get_columns("projects")}
    conversation_columns = {item["name"] for item in inspect(legacy).get_columns("conversations")}
    usage_columns = {item["name"] for item in inspect(legacy).get_columns("llm_usage_logs")}
    assert "analysis_brief" in project_columns
    assert "deleted_at" in {item["name"] for item in inspect(legacy).get_columns("reports")}
    assert "clarification_mode" in conversation_columns
    assert {"stage", "run_id", "report_id", "dashboard_id", "purpose", "usage_unavailable"} <= usage_columns
    legacy.dispose()
