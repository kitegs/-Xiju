from io import BytesIO
import asyncio
import hashlib
import time
import pytest
from datetime import datetime, timedelta, timezone

from docx import Document
from fastapi.testclient import TestClient
from uuid import uuid4

from app.main import app


def test_token_usage_center_groups_provider_usage_without_currency_fields():
    from app.database import SessionLocal
    from app.models import LlmUsageLog

    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()

        async def seed_usage():
            async with SessionLocal() as session:
                session.add_all([
                    LlmUsageLog(
                        workspace_id=workspace["id"], provider="deepseek", model="deepseek-test",
                        stage="planning", purpose="生成分析计划", prompt_tokens=100,
                        cache_hit_tokens=20, cache_miss_tokens=80, completion_tokens=30,
                        total_tokens=130, latency_ms=50,
                    ),
                    LlmUsageLog(
                        workspace_id=workspace["id"], provider="deepseek", model="deepseek-test",
                        stage="synthesis", purpose="整理运行证据结论", run_id="run-1",
                        report_id="report-1", usage_unavailable=True, latency_ms=60,
                    ),
                ])
                await session.commit()

        asyncio.run(seed_usage())
        response = client.get("/api/v1/usage/tokens", params={
            "workspace_id": workspace["id"], "range": "all", "group_by": "stage",
        })
        assert response.status_code == 200
        summary = response.json()
        assert summary["totals"] == {
            "prompt_tokens": 100, "cache_hit_tokens": 20, "cache_miss_tokens": 80,
            "completion_tokens": 30, "total_tokens": 130, "call_count": 2,
        }
        assert sum(group["total_tokens"] for group in summary["groups"]) == 130
        assert {group["key"] for group in summary["groups"]} == {"planning", "synthesis"}
        assert summary["unavailable_call_count"] == 1
        assert all("estimated_cost_cny" not in row and "pricing" not in row for row in summary["latest"])

        run_summary = client.get("/api/v1/usage/tokens/runs/run-1", params={
            "workspace_id": workspace["id"],
        }).json()
        assert run_summary["totals"]["call_count"] == 1
        assert run_summary["unavailable_call_count"] == 1

        report_summary = client.get("/api/v1/usage/tokens/reports/report-1", params={
            "workspace_id": workspace["id"],
        }).json()
        assert report_summary["totals"]["call_count"] == 1


def test_token_usage_center_filters_today_seven_days_and_all():
    from app.database import SessionLocal
    from app.models import LlmUsageLog

    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()

        async def seed_usage():
            async with SessionLocal() as session:
                session.add_all([
                    LlmUsageLog(
                        workspace_id=workspace["id"], provider="deepseek", model="recent",
                        stage="planning", prompt_tokens=10, completion_tokens=5, total_tokens=15,
                        created_at=datetime.now(timezone.utc),
                    ),
                    LlmUsageLog(
                        workspace_id=workspace["id"], provider="deepseek", model="old",
                        stage="chart", prompt_tokens=20, completion_tokens=10, total_tokens=30,
                        created_at=datetime.now(timezone.utc) - timedelta(days=8),
                    ),
                ])
                await session.commit()

        asyncio.run(seed_usage())
        today = client.get("/api/v1/usage/tokens", params={
            "workspace_id": workspace["id"], "range": "today", "group_by": "stage",
        }).json()
        seven_days = client.get("/api/v1/usage/tokens", params={
            "workspace_id": workspace["id"], "range": "7d", "group_by": "stage",
        }).json()
        all_time = client.get("/api/v1/usage/tokens", params={
            "workspace_id": workspace["id"], "range": "all", "group_by": "stage",
        }).json()
        assert today["totals"]["total_tokens"] == 15
        assert seven_days["totals"]["total_tokens"] == 15
        assert all_time["totals"]["total_tokens"] == 45
        for summary in (today, seven_days, all_time):
            assert sum(group["total_tokens"] for group in summary["groups"]) == summary["totals"]["total_tokens"]


@pytest.mark.parametrize('reject_summary', [False, True, "bound"])
def test_model_planning_and_synthesis_usage_are_recorded_separately(monkeypatch, reject_summary):
    import app.main as main_module
    from app.llm import CompletionResult

    async def fake_complete(config, messages, context, request_options=None):
        planning = messages[0]["content"].startswith("[planner:")
        content = '{"steps": []}' if planning else ('{"summary":"未绑定的数字 123","claims":[],"limitations":[]}' if reject_summary is True else '{"summary":"已按证据完成分析。","claims":[],"limitations":[]}')
        if not planning and reject_summary == "bound":
            import json
            packet = json.loads(context)
            content = json.dumps({'protocol_version': 'synthesis-v5', 'summary': '已核对事实。',
                                  'facts': [packet['facts'][0]['fact_id']]})
        elif not planning and reject_summary is False:
            import json
            packet = json.loads(context)
            content = json.dumps({'protocol_version': 'synthesis-v5', 'summary': '已核对事实。',
                                  'facts': [packet['facts'][0]['fact_id']]})
        return CompletionResult(
            content=content,
            request_id="planning-request" if planning else "synthesis-request",
            model=config.model,
            usage={
                "prompt_tokens": 10, "cache_hit_tokens": 0, "cache_miss_tokens": 10,
                "completion_tokens": 5, "total_tokens": 15,
            },
            estimated_cost_cny=0, latency_ms=1, pricing={}, usage_available=True,
        )

    monkeypatch.setattr(main_module, "complete", fake_complete)
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()
        dataset = client.post("/api/v1/samples/retail-sales-2026/import", params={
            "workspace_id": workspace["id"],
        }).json()
        provider = client.put("/api/v1/settings/providers/deepseek", params={
            "workspace_id": workspace["id"],
        }, json={
            "enabled": True, "is_default": True, "base_url": "https://api.deepseek.com",
            "model": "deepseek-test", "api_key": "test-key", "options": {},
        })
        assert provider.status_code == 200
        conversation = client.post("/api/v1/conversations", json={
            "workspace_id": workspace["id"], "dataset_id": dataset["id"],
            "title": "Token 阶段测试", "clarification_mode": "off",
        }).json()
        planned = client.post("/api/v1/chat/plan", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "dataset_id": dataset["id"], "message": "分析这份数据",
        })
        assert planned.status_code == 200
        executed = client.post("/api/v1/chat/execute", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "plan_message_id": planned.json()["user_message"]["id"], "approved": True,
        })
        assert executed.status_code == 200
        assert executed.json()["provider"] == ("deepseek-fallback" if reject_summary is True else "deepseek"), executed.json()
        if reject_summary == "bound":
            assert "{{" not in executed.json()["assistant_message"]["content"]
            assert executed.json()["assistant_message"]["message_meta"]["claims"]["claims"][0]["evidence_ids"]
        usage = client.get("/api/v1/usage/tokens", params={
            "workspace_id": workspace["id"], "range": "all", "group_by": "stage",
        }).json()
        assert {group["key"] for group in usage["groups"]} == {"planning", "synthesis"}
        assert usage["totals"]["total_tokens"] == 30


def test_project_analysis_brief_and_conversation_clarification_mode_are_persistent():
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()
        dataset = client.post("/api/v1/samples/retail-sales-2026/import", params={
            "workspace_id": workspace["id"],
        }).json()
        project = client.get("/api/v1/projects", params={"workspace_id": workspace["id"]}).json()[0]
        conversation = client.post("/api/v1/conversations", json={
            "workspace_id": workspace["id"], "dataset_id": dataset["id"], "title": "澄清测试",
        }).json()
        assert conversation["clarification_mode"] == "auto"

        questions = client.post(f"/api/v1/projects/{project['id']}/brief/questions", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "dataset_id": dataset["id"], "objective": "生成经营报告",
        })
        assert questions.status_code == 200
        decision = questions.json()
        assert decision["should_pause"] is True
        assert len(decision["questions"]) == 8

        saved = client.put(f"/api/v1/projects/{project['id']}/brief", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "dataset_id": dataset["id"], "objective": "生成经营报告",
            "answers": {"audience": "总经理"}, "use_recommended_defaults": True,
        })
        assert saved.status_code == 200
        brief = saved.json()
        assert brief["answers"]["audience"] == "总经理"
        assert brief["confirmed_keys"] == ["audience"]
        assert len(brief["defaulted_keys"]) == 7

        repeated = client.post(f"/api/v1/projects/{project['id']}/brief/questions", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "dataset_id": dataset["id"], "objective": "生成经营报告", "mode": "always",
        }).json()
        assert repeated["should_pause"] is False
        assert repeated["questions"] == []

        changed = client.patch(f"/api/v1/conversations/{conversation['id']}", json={
            "clarification_mode": "off",
        })
        assert changed.status_code == 200
        assert changed.json()["clarification_mode"] == "off"
        listed = client.get("/api/v1/conversations", params={"workspace_id": workspace["id"]}).json()
        assert listed[0]["clarification_mode"] == "off"


def test_chat_plan_pauses_for_intake_then_continues_with_visible_assumptions():
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()
        dataset = client.post("/api/v1/samples/retail-sales-2026/import", params={
            "workspace_id": workspace["id"],
        }).json()
        project = client.get("/api/v1/projects", params={"workspace_id": workspace["id"]}).json()[0]
        conversation = client.post("/api/v1/conversations", json={
            "workspace_id": workspace["id"], "dataset_id": dataset["id"], "title": "新分析",
        }).json()
        request = {
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "dataset_id": dataset["id"], "message": "生成管理层报告", "analysis_mode": "business",
        }
        paused = client.post("/api/v1/chat/plan", json=request)
        assert paused.status_code == 200
        assert paused.json()["intake"]["should_pause"] is True
        assert paused.json()["plan"]["blocked"] is True
        assert paused.json()["plan"]["steps"] == []

        accepted = client.put(f"/api/v1/projects/{project['id']}/brief", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "dataset_id": dataset["id"], "objective": request["message"],
            "answers": {"audience": "管理层"}, "use_recommended_defaults": True,
        })
        assert accepted.status_code == 200
        planned = client.post("/api/v1/chat/plan", json={
            **request, "resume_message_id": paused.json()["user_message"]["id"],
        })
        assert planned.status_code == 200
        assert planned.json()["intake"] is None
        assert planned.json()["user_message"]["message_meta"]["intake_completed"] is True
        assert any(step["tool"] == "report.build" for step in planned.json()["plan"]["steps"])
        assert len(client.get(f"/api/v1/conversations/{conversation['id']}/messages").json()) == 1

        executed = client.post("/api/v1/chat/execute", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "plan_message_id": planned.json()["user_message"]["id"], "approved": True,
        })
        assert executed.status_code == 200
        report_id = executed.json()["report_id"]
        report = client.get(f"/api/v1/reports/{report_id}").json()
        assumptions = report["document"]["metadata"]["unconfirmed_assumptions"]
        assert len(assumptions) == 7
        assert any(block["id"] == "unconfirmed-assumptions" for block in report["document"]["blocks"])


def test_local_queue_records_completed_runs_and_history():
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()
        dataset = client.post("/api/v1/samples/retail-sales-2026/import", params={"workspace_id": workspace["id"]}).json()
        conversation = client.post("/api/v1/conversations", json={"workspace_id": workspace["id"], "dataset_id": dataset["id"], "title": "队列"}).json()
        planned = client.post("/api/v1/chat/plan", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"], "dataset_id": dataset["id"],
            "analysis_mode": "research", "message": "对数据做科研统计",
        }).json()
        queued = client.post("/api/v1/chat/execute-async", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"], "plan_message_id": planned["user_message"]["id"],
        })
        assert queued.status_code == 202
        run = queued.json(); assert run["status"] in {"queued", "running"}
        for _ in range(80):
            current = client.get("/api/v1/analysis-runs", params={"workspace_id": workspace["id"], "conversation_id": conversation["id"]}).json()[0]
            if current["status"] not in {"queued", "running"}:
                break
            time.sleep(0.05)
        assert current["status"] == "completed"
        assert current["result_message_id"]
        repeated = client.post("/api/v1/chat/execute-async", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "plan_message_id": planned["user_message"]["id"],
        })
        assert repeated.status_code == 202, repeated.text
        assert repeated.json()["id"] == current["id"]
        events = client.get(f"/api/v1/analysis-runs/{current['id']}/events", params={"workspace_id": workspace["id"]}).json()
        stream = client.get(f"/api/v1/analysis-runs/{current['id']}/events/stream", params={"workspace_id": workspace["id"]})
        assert stream.status_code == 200
        assert "text/event-stream" in stream.headers["content-type"]
        assert f"id: {events[-1]['sequence']}" in stream.text
        tail = client.get(f"/api/v1/analysis-runs/{current['id']}/events/stream", params={"workspace_id": workspace["id"], "after": events[-1]["sequence"]})
        assert tail.text == ""


def test_running_task_cancel_propagates_and_persists_terminal_event(monkeypatch):
    from app import main as main_module

    async def slow_execution(*args, **kwargs):
        await asyncio.sleep(30)

    monkeypatch.setattr(main_module, "_execute_chat_plan", slow_execution)
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()
        conversation = client.post("/api/v1/conversations", json={
            "workspace_id": workspace["id"], "dataset_id": None, "title": "取消测试",
        }).json()
        planned = client.post("/api/v1/chat/plan", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "message": "说明如何开始", "execution_mode": "safe",
        }).json()
        queued = client.post("/api/v1/chat/execute-async", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "plan_message_id": planned["user_message"]["id"],
        }).json()
        for _ in range(50):
            current = client.get("/api/v1/analysis-runs", params={
                "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            }).json()[0]
            if current["status"] == "running":
                break
            time.sleep(0.02)
        assert current["status"] == "running"
        assert client.post(f"/api/v1/analysis-runs/{queued['id']}/cancel", params={"workspace_id": workspace["id"]}).status_code == 200
        for _ in range(50):
            current = client.get("/api/v1/analysis-runs", params={
                "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            }).json()[0]
            if current["status"] == "cancelled":
                break
            time.sleep(0.02)
        assert current["status"] == "cancelled"
        events = client.get(f"/api/v1/analysis-runs/{queued['id']}/events", params={"workspace_id": workspace["id"]}).json()
        assert events[-1]["event_type"] == "run.cancelled"


def test_startup_recovery_marks_abandoned_running_task_interrupted(monkeypatch):
    from app import main as main_module
    from app.models import AnalysisRun
    from app.database import SessionLocal

    monkeypatch.setattr(main_module, "_schedule_run", lambda *args, **kwargs: None)
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()
        conversation = client.post("/api/v1/conversations", json={
            "workspace_id": workspace["id"], "dataset_id": None, "title": "恢复测试",
        }).json()
        planned = client.post("/api/v1/chat/plan", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"], "message": "说明如何开始",
        }).json()
        queued = client.post("/api/v1/chat/execute-async", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "plan_message_id": planned["user_message"]["id"],
        }).json()

        async def abandon_and_recover():
            async with SessionLocal() as session:
                run = await session.get(AnalysisRun, queued["id"])
                run.status = "running"
                await session.commit()
            await main_module._recover_local_runs()

        asyncio.run(abandon_and_recover())
        recovered = client.get("/api/v1/analysis-runs", params={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
        }).json()[0]
        assert recovered["status"] == "interrupted"
        events = client.get(f"/api/v1/analysis-runs/{queued['id']}/events", params={"workspace_id": workspace["id"]}).json()
        assert events[-1]["event_type"] == "interrupted"


def test_queued_recovery_and_retry_are_idempotent(monkeypatch):
    from app import main as main_module
    scheduled = []
    monkeypatch.setattr(main_module, '_schedule_run', lambda run_id, *args: scheduled.append(run_id))
    with TestClient(app) as client:
        wid = client.post('/api/v1/workspaces/bootstrap').json()['id']
        conversation = client.post('/api/v1/conversations', json={'workspace_id': wid, 'title': '恢复排队'}).json()
        planned = client.post('/api/v1/chat/plan', json={'workspace_id': wid, 'conversation_id': conversation['id'], 'message': '如何开始'}).json()
        payload = {'workspace_id': wid, 'conversation_id': conversation['id'], 'plan_message_id': planned['user_message']['id']}
        queued = client.post('/api/v1/chat/execute-async', json=payload).json()
        scheduled.clear()
        asyncio.run(main_module._recover_local_runs())
        assert scheduled == [queued['id']]
        client.post(f"/api/v1/analysis-runs/{queued['id']}/cancel", params={'workspace_id': wid})
        first = client.post(f"/api/v1/analysis-runs/{queued['id']}/retry", params={'workspace_id': wid})
        second = client.post(f"/api/v1/analysis-runs/{queued['id']}/retry", params={'workspace_id': wid})
        assert first.status_code == second.status_code == 202
        assert first.json()['id'] == second.json()['id']
        assert first.json()['attempt'] == 2
        assert scheduled.count(first.json()['id']) == 1


def test_sample_to_chat_to_report_and_settings_keeps_keys_masked():
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()
        dataset = client.post("/api/v1/samples/retail-sales-2026/import", params={"workspace_id": workspace["id"]}).json()
        preview = client.get(f"/api/v1/datasets/{dataset['id']}/preview", params={"workspace_id": workspace["id"], "limit": 5})
        assert preview.status_code == 200
        assert len(preview.json()["rows"]) == 5
        assert preview.json()["quality"]["score"] == 100

        clean_preview = client.post(f"/api/v1/datasets/{dataset['id']}/cleaning/preview", json={
            "workspace_id": workspace["id"], "steps": [{"operation": "filter_rows", "column": "销售额", "operator": "gte", "value": 60000}],
        })
        assert clean_preview.status_code == 200
        assert clean_preview.json()["after_rows"] < clean_preview.json()["before_rows"]

        clean_apply = client.post(f"/api/v1/datasets/{dataset['id']}/cleaning/apply", json={
            "workspace_id": workspace["id"], "name": "高销售记录", "steps": [{"operation": "filter_rows", "column": "销售额", "operator": "gte", "value": 60000}],
        })
        assert clean_apply.status_code == 201
        assert clean_apply.json()["dataset"]["id"] != dataset["id"]
        assert clean_apply.json()["recipe"]["source_dataset_id"] == dataset["id"]
        assert clean_apply.json()["recipe"]["result_dataset_id"] == clean_apply.json()["dataset"]["id"]
        assert clean_apply.json()["recipe"]["impact"]["changed_cells"] == 0
        original_after_cleaning = client.get(f"/api/v1/datasets/{dataset['id']}/preview", params={"workspace_id": workspace["id"], "limit": 5})
        assert original_after_cleaning.json()["total_rows"] == dataset["profile"]["row_count"]
        conversation = client.post("/api/v1/conversations", json={"workspace_id": workspace["id"], "dataset_id": dataset["id"], "title": "新分析"}).json()

        result = client.post("/api/v1/chat", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"], "dataset_id": dataset["id"],
            "message": "生成一份管理层报告",
        })
        assert result.status_code == 200
        assert result.json()["report_id"]
        assert len(client.get(f"/api/v1/conversations/{conversation['id']}/messages").json()) == 2

        preview = client.post(f"/api/v1/datasets/{dataset['id']}/chart-preview", json={
            "workspace_id": workspace["id"], "x_column": "区域", "y_column": "销售额", "aggregate": "sum",
        })
        assert preview.status_code == 200
        assert preview.json()["data"][0]["区域"]

        saved = client.put("/api/v1/settings/providers/openai", params={"workspace_id": workspace["id"]}, json={
            "enabled": False, "is_default": True, "base_url": "https://api.openai.com/v1", "model": "gpt-5.6-terra",
            "api_key": "sk-test-not-a-real-key", "options": {"reasoning_effort": "medium"},
        })
        assert saved.status_code == 200
        assert saved.json()["has_api_key"] is True
        assert "sk-test-not-a-real-key" not in saved.text

        settings = client.put("/api/v1/settings/app", params={"workspace_id": workspace["id"]}, json={
            "language": "zh-CN", "theme": "light", "autosave": True, "autosave_interval_seconds": 30,
            "safe_mode": True, "telemetry": False, "default_export": "pdf", "confirm_external_requests": True,
            "default_clarification_mode": "always",
        })
        assert settings.status_code == 200
        assert settings.json()["safe_mode"] is True
        inherited = client.post("/api/v1/conversations", json={
            "workspace_id": workspace["id"], "title": "继承澄清设置",
        })
        assert inherited.status_code == 201
        assert inherited.json()["clarification_mode"] == "always"

        unsafe_settings = client.put("/api/v1/settings/app", params={"workspace_id": workspace["id"]}, json={
            "language": "zh-CN", "theme": "light", "autosave": True, "autosave_interval_seconds": 30,
            "safe_mode": False, "telemetry": False, "default_export": "pdf", "confirm_external_requests": True,
        })
        assert unsafe_settings.status_code == 200
        assert unsafe_settings.json()["safe_mode"] is True
        assert client.get("/api/v1/settings/app", params={"workspace_id": workspace["id"]}).json()["safe_mode"] is True


def test_safe_readonly_plan_runs_without_approval_and_persists_tool_evidence():
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()
        dataset = client.post(
            "/api/v1/samples/retail-sales-2026/import", params={"workspace_id": workspace["id"]}
        ).json()
        conversation = client.post("/api/v1/conversations", json={
            "workspace_id": workspace["id"], "dataset_id": dataset["id"], "title": "新分析",
            "clarification_mode": "off",
        }).json()

        planned = client.post("/api/v1/chat/plan", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "dataset_id": dataset["id"], "message": "生成一份管理层报告并说明数据质量",
        })
        assert planned.status_code == 200
        planned_body = planned.json()
        assert planned_body["plan"]["status"] == "pending"
        assert planned_body["plan"]["requires_approval"] is False
        assert planned_body["plan"]["blocked"] is False
        assert {item["decision"] for item in planned_body["plan"]["policy"]} == {"allow"}
        assert [step["tool"] for step in planned_body["plan"]["steps"]] == [
            "dataset.profile", "data.quality", "report.build", "report.layout", "report.review", "assistant.synthesize",
        ]
        assert len(client.get(f"/api/v1/conversations/{conversation['id']}/messages").json()) == 1

        executed = client.post("/api/v1/chat/execute", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "plan_message_id": planned_body["user_message"]["id"],
        })
        assert executed.status_code == 200
        executed_body = executed.json()
        assert executed_body["report_id"]
        assert executed_body["user_message"]["message_meta"]["analysis_plan"]["status"] == "completed"
        tool_runs = executed_body["assistant_message"]["message_meta"]["tool_runs"]
        assert len(tool_runs) == 6
        assert all(run["status"] == "completed" and run["duration_ms"] >= 1 for run in tool_runs)
        evidence = executed_body["assistant_message"]["message_meta"]["evidence"]
        assert {item["id"] for item in evidence} >= {"dataset-shape", "quality-summary"}

        messages = client.get(f"/api/v1/conversations/{conversation['id']}/messages").json()
        assert len(messages) == 2
        assert messages[0]["message_meta"]["analysis_plan"]["status"] == "completed"
        assert messages[1]["message_meta"]["tool_runs"][0]["tool"] == "dataset.profile"


def test_pending_analysis_plan_can_be_cancelled():
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()
        conversation = client.post("/api/v1/conversations", json={
            "workspace_id": workspace["id"], "dataset_id": None, "title": "新分析",
        }).json()
        planned = client.post("/api/v1/chat/plan", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "dataset_id": None, "message": "先告诉我如何开始",
        }).json()
        cancelled = client.post(
            f"/api/v1/chat/plans/{planned['user_message']['id']}/cancel",
            params={"workspace_id": workspace["id"], "conversation_id": conversation["id"]},
        )
        assert cancelled.status_code == 200
        assert cancelled.json()["message_meta"]["analysis_plan"]["status"] == "cancelled"
        rejected = client.post("/api/v1/chat/execute", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "plan_message_id": planned["user_message"]["id"],
        })
        assert rejected.status_code == 409


def test_multiseries_chart_sql_evidence_and_report_version_restore():
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()
        dataset = client.post(
            "/api/v1/samples/retail-sales-2026/import", params={"workspace_id": workspace["id"]}
        ).json()
        preview = client.post(f"/api/v1/datasets/{dataset['id']}/chart-preview", json={
            "workspace_id": workspace["id"], "x_column": "区域", "y_column": "销售额",
            "series_columns": ["成本"], "aggregate": "sum", "limit": 10,
        })
        assert preview.status_code == 200
        assert {"区域", "销售额", "成本"} <= set(preview.json()["data"][0])

        sql = client.post(f"/api/v1/datasets/{dataset['id']}/sql", json={
            "workspace_id": workspace["id"],
            "sql": 'SELECT "区域", SUM("销售额") AS "销售额" FROM dataset GROUP BY 1 ORDER BY 2 DESC',
            "max_rows": 10,
        })
        assert sql.status_code == 200
        assert sql.json()["evidence"]["code"].startswith("SELECT")

        analysis = client.post("/api/v1/analysis/run", json={
            "workspace_id": workspace["id"], "dataset_id": dataset["id"], "prompt": "生成报告",
        }).json()
        report = client.get(f"/api/v1/reports/{analysis['report_id']}").json()
        original_title = report["title"]
        report["title"] = "人工修改标题"
        saved = client.put(f"/api/v1/reports/{report['id']}", json={"title": report["title"], "document": report["document"]})
        assert saved.status_code == 200
        versions = client.get(f"/api/v1/reports/{report['id']}/versions").json()
        assert versions[0]["title"] == original_title
        restored = client.post(f"/api/v1/reports/{report['id']}/versions/{versions[0]['id']}/restore")
        assert restored.status_code == 200
        assert restored.json()["title"] == original_title
        quality = client.get(f"/api/v1/reports/{report['id']}/quality")
        assert quality.status_code == 200
        assert quality.json()["passed"] is True and quality.json()["score"] >= 70
        exported = client.get(f"/api/v1/reports/{report['id']}/export.docx")
        assert exported.status_code == 200
        assert exported.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert exported.content[:2] == b"PK"


def test_chart_patch_recomputes_full_data_requires_approval_and_keeps_version_history():
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()
        dataset = client.post("/api/v1/samples/retail-sales-2026/import", params={"workspace_id": workspace["id"]}).json()
        report_id = client.post("/api/v1/analysis/run", json={
            "workspace_id": workspace["id"], "dataset_id": dataset["id"], "prompt": "生成报告",
        }).json()["report_id"]
        report = client.get(f"/api/v1/reports/{report_id}").json()
        chart_id = next(chart["id"] for chart in report["document"]["charts"] if chart["chart_type"] == "bar")

        context = client.get(f"/api/v1/reports/{report_id}/charts/{chart_id}/context")
        assert context.status_code == 200
        base_version = context.json()["report_version_hint"]
        proposal = client.post(f"/api/v1/reports/{report_id}/charts/{chart_id}/proposals", json={
            "workspace_id": workspace["id"], "instruction": "改成横向条形图，只显示前 2 项，使用商务配色并显示标签",
        })
        assert proposal.status_code == 200
        assert proposal.json()["patch"]["chart_type"] == "bar"
        assert proposal.json()["patch"]["limit"] == 2

        preview_payload = {"workspace_id": workspace["id"], "base_version": base_version, "patch": proposal.json()["patch"]}
        preview = client.post(f"/api/v1/reports/{report_id}/charts/{chart_id}/preview", json=preview_payload)
        assert preview.status_code == 200
        assert len(preview.json()["data"]) == 2
        assert preview.json()["evidence"]["source_columns"]
        assert preview.json()["chart"]["data"] == preview.json()["data"]

        saved = client.post(f"/api/v1/reports/{report_id}/charts/{chart_id}/apply", json={**preview_payload, "approved": True})
        assert saved.status_code == 200
        changed = next(chart for chart in saved.json()["document"]["charts"] if chart["id"] == chart_id)
        assert changed["style"]["orientation"] == "horizontal"
        assert len(changed["data"]) == 2
        assert client.get(f"/api/v1/reports/{report_id}/versions").json()


def test_team_roles_review_comments_share_and_audit():
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()
        workspace_id = workspace["id"]
        context = client.get("/api/v1/team/context", params={"workspace_id": workspace_id})
        assert context.status_code == 200
        assert context.json()["current_role"] == "owner"

        invited = client.post("/api/v1/team/members", params={"workspace_id": workspace_id}, json={
            "email": f"analyst-{uuid4().hex[:8]}@example.com", "display_name": "数据分析师", "role": "analyst",
        })
        assert invited.status_code == 201
        assert invited.json()["role"] == "analyst"
        analyst_headers = {"X-Actor-ID": invited.json()["user_id"]}

        dataset = client.post("/api/v1/samples/retail-sales-2026/import", params={"workspace_id": workspace_id}).json()
        analysis = client.post("/api/v1/analysis/run", json={
            "workspace_id": workspace_id, "dataset_id": dataset["id"], "prompt": "生成报告",
        }).json()
        report_id = analysis["report_id"]
        denied_publish = client.post(f"/api/v1/reports/{report_id}/workflow/publish", headers=analyst_headers)
        assert denied_publish.status_code == 403
        assert client.get("/api/v1/team/context", params={"workspace_id": workspace_id}, headers={"X-Actor-ID": "not-a-member"}).status_code == 403

        comment = client.post(f"/api/v1/reports/{report_id}/comments", json={"body": "请核对利润口径"})
        assert comment.status_code == 201
        assert client.patch(f"/api/v1/reports/{report_id}/comments/{comment.json()['id']}/resolve").json()["resolved"] is True

        assert client.post(f"/api/v1/reports/{report_id}/workflow/submit").json()["status"] == "in_review"
        assert client.post(f"/api/v1/reports/{report_id}/workflow/approve").json()["status"] == "approved"
        assert client.post(f"/api/v1/reports/{report_id}/workflow/publish").json()["status"] == "published"
        share = client.post(f"/api/v1/reports/{report_id}/share", json={"expires_in_hours": 24})
        assert share.status_code == 201
        token = share.json()["url"].rsplit("/", 1)[-1]
        assert client.get(f"/api/v1/shared/reports/{token}").json()["id"] == report_id
        audit = client.get("/api/v1/audit", params={"workspace_id": workspace_id})
        assert audit.status_code == 200
        assert {item["action"] for item in audit.json()} >= {"member.invite", "report.publish", "share.create"}


def test_superset_dashboard_mapping_is_idempotent(monkeypatch):
    from app import main as main_module

    async def fake_create(title, allowed_domains):
        return {"dashboard_id": 42, "embedded_id": f"embedded-{uuid4().hex}", "title": title}

    monkeypatch.setattr(main_module, "create_embedded_dashboard", fake_create)
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces", json={"name": f"Superset {uuid4().hex[:8]}", "description": "test"}).json()
        url = f"/api/v1/integrations/superset/bootstrap-dashboard?workspace_id={workspace['id']}"
        first = client.post(url, json={"title": "专业看板", "allowed_domains": ["http://localhost:5175"]})
        second = client.post(url, json={"title": "不应重复创建"})
        assert first.status_code == 200
        assert second.json()["embedded_id"] == first.json()["embedded_id"]
        resources = client.get("/api/v1/integrations/superset/resources", params={"workspace_id": workspace["id"]}).json()
        assert len(resources) == 1
        assert resources[0]["dashboard_id"] == 42


def test_analysis_modes_route_plans_to_scenario_toolkits():
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()
        dataset = client.post("/api/v1/samples/retail-sales-2026/import", params={"workspace_id": workspace["id"]}).json()
        modes = client.get("/api/v1/analysis-modes").json()
        assert modes["default"] == "auto"
        assert {item["id"] for item in modes["modes"]} >= {"business", "research", "survey", "professional"}
        conversation = client.post("/api/v1/conversations", json={
            "workspace_id": workspace["id"], "dataset_id": dataset["id"], "title": "新分析",
        }).json()
        planned = client.post("/api/v1/chat/plan", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"], "dataset_id": dataset["id"],
            "analysis_mode": "research", "message": "给出描述统计和置信区间",
        }).json()
        assert planned["plan"]["analysis_mode"] == "research"
        assert planned["plan"]["mode_reason"] == "用户手动选择"
        assert "statistics.describe" in [step["tool"] for step in planned["plan"]["steps"]]


def test_professional_publish_is_blocked_or_requires_explicit_approval_by_mode():
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()
        dataset = client.post("/api/v1/samples/retail-sales-2026/import", params={"workspace_id": workspace["id"]}).json()
        conversation = client.post("/api/v1/conversations", json={
            "workspace_id": workspace["id"], "dataset_id": dataset["id"], "title": "权限矩阵",
        }).json()
        for mode, expected_status, blocked in (("safe", 403, True), ("partial", 409, False), ("full", 409, False)):
            planned = client.post("/api/v1/chat/plan", json={
                "workspace_id": workspace["id"], "conversation_id": conversation["id"],
                "dataset_id": dataset["id"], "analysis_mode": "professional",
                "execution_mode": mode, "message": "生成专业看板",
            }).json()
            assert planned["plan"]["blocked"] is blocked
            publish_policy = next(item for item in planned["plan"]["policy"] if item["tool"] == "superset.publish")
            assert publish_policy["decision"] == ("forbid" if mode == "safe" else "approval")
            rejected = client.post("/api/v1/chat/execute-async", json={
                "workspace_id": workspace["id"], "conversation_id": conversation["id"],
                "plan_message_id": planned["user_message"]["id"], "approved": False,
            })
            assert rejected.status_code == expected_status


def test_superset_publish_preview_and_resource_mapping(monkeypatch):
    from app import main as main_module

    async def fake_publish(**kwargs):
        return {
            "database_id": 7, "superset_dataset_id": 8, "dashboard_id": 9,
            "embedded_id": "embedded-published", "title": kwargs["title"],
            "charts": [{"id": 10, "uuid": "chart-uuid", "title": "销售额", "chart_spec_id": "kpi-sales", "viz_type": "big_number_total"}],
            "dashboard_url": "http://localhost:8088/superset/dashboard/9/",
            "table_name": "dataset_test", "dataset_name": kwargs["dataset_name"],
        }

    monkeypatch.setattr(main_module, "publish_superset_dashboard", fake_publish)
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces", json={"name": f"Publish {uuid4().hex[:8]}", "description": "test"}).json()
        dataset = client.post("/api/v1/samples/retail-sales-2026/import", params={"workspace_id": workspace["id"]}).json()
        preview = client.post("/api/v1/integrations/superset/publish-preview", json={
            "workspace_id": workspace["id"], "dataset_id": dataset["id"],
        })
        assert preview.status_code == 200
        assert preview.json()["requires_approval"] is True
        assert preview.json()["chart_count"] >= 1
        published = client.post("/api/v1/integrations/superset/publish", json={
            "workspace_id": workspace["id"], "dataset_id": dataset["id"], "title": "零售专业看板",
        })
        assert published.status_code == 200
        assert published.json()["dashboard_id"] == 9
        resources = client.get("/api/v1/integrations/superset/resources", params={"workspace_id": workspace["id"]}).json()
        assert {item["resource_type"] for item in resources} >= {"database", "dataset", "chart", "dashboard"}


def test_forecast_options_and_template_upload_are_auditable():
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces/bootstrap").json()
        dataset = client.post("/api/v1/samples/retail-sales-2026/import", params={"workspace_id": workspace["id"]}).json()
        conversation = client.post("/api/v1/conversations", json={"workspace_id": workspace["id"], "dataset_id": dataset["id"], "title": "预测"}).json()
        planned = client.post("/api/v1/chat/plan", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"], "dataset_id": dataset["id"],
            "message": "预测未来销售并给我行动建议", "analysis_options": {"include_recommendations": True, "forecast": {"enabled": True, "date_column": "日期", "target_column": "销售额", "horizon": 3, "frequency": "MS", "aggregate": "sum"}},
        }).json()
        assert {"timeseries.forecast", "action.recommend"} <= {step["tool"] for step in planned["plan"]["steps"]}
        result = client.post("/api/v1/chat/execute", json={"workspace_id": workspace["id"], "conversation_id": conversation["id"], "plan_message_id": planned["user_message"]["id"]})
        assert result.status_code == 200
        meta = result.json()["assistant_message"]["message_meta"]
        assert meta["recommendations"] and any(run["tool"] == "timeseries.forecast" and run["code"] for run in meta["tool_runs"])
        word = Document(); word.add_paragraph("{{ report.title }}"); word.add_paragraph("{{ report.summary }}")
        data = BytesIO(); word.save(data)
        uploaded = client.post("/api/v1/report-templates/upload", data={"workspace_id": workspace["id"], "name": "固定格式"}, files={"file": ("fixed.docx", data.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        assert uploaded.status_code == 201
        report_id = client.post("/api/v1/analysis/run", json={"workspace_id": workspace["id"], "dataset_id": dataset["id"]}).json()["report_id"]
        rendered = client.post(f"/api/v1/report-templates/{uploaded.json()['id']}/render", json={"workspace_id": workspace["id"], "report_id": report_id})
        assert rendered.status_code == 201
        assert client.get(rendered.json()["download_url"]).status_code == 200


def _global_superstore_csv() -> bytes:
    columns = [
        "Category", "City", "Country", "Customer ID", "Customer Name", "Discount", "Market",
        "ji_lu-shu", "Order Date", "Order ID", "Order Priority", "Product ID", "Product Name",
        "Profit", "Quantity", "Region", "Row ID", "Sales", "Segment", "Ship Date", "Ship Mode",
        "Shipping Cost", "State", "Sub-Category", "Year", "Market2", "weeknum",
    ]
    rows = [",".join(columns)]
    regions = ("West", "East", "Central", "South")
    categories = ("Technology", "Furniture", "Office Supplies")
    for index in range(51_290):
        year = 2011 + index % 4
        month = 1 + index % 12
        day = 1 + index % 27
        sales = 100 + index % 1_900
        profit = (index % 401) - 120
        row = [
            categories[index % 3], f"City{index % 80}", f"Country{index % 12}", f"C{index % 2000:04d}",
            f"Customer{index % 2000}", f"{(index % 6) / 10:.1f}", f"Market{index % 5}", str(index + 1),
            f"{year}-{month:02d}-{day:02d}", f"O{index + 1:06d}", ("Low", "Medium", "High")[index % 3],
            f"P{index % 500:04d}", f"Product{index % 500}", str(profit), str(1 + index % 8),
            regions[index % 4], str(index + 1), str(sales), ("Consumer", "Corporate", "Home Office")[index % 3],
            f"{year}-{month:02d}-{min(day + 2, 28):02d}", ("First Class", "Standard Class")[index % 2],
            f"{5 + index % 60}.0", f"State{index % 30}", f"Sub{index % 9}", str(year),
            f"Market{index % 5}", str(1 + index % 52),
        ]
        rows.append(",".join(row))
    return ("\n".join(rows) + "\n").encode("utf-8")


def test_global_superstore_project_lineage_is_repeatable_end_to_end():
    content = _global_superstore_csv()
    expected_hash = hashlib.sha256(content).hexdigest()
    with TestClient(app) as client:
        workspace = client.post("/api/v1/workspaces", json={
            "name": f"Global Superstore {uuid4().hex[:8]}", "description": "repeatable e2e",
        }).json()
        uploaded = client.post(
            "/api/v1/datasets/upload", data={"workspace_id": workspace["id"]},
            files={"file": ("Global Superstore.csv", content, "text/csv")},
        )
        assert uploaded.status_code == 201
        dataset = uploaded.json()
        assert dataset["profile"]["row_count"] == 51_290
        assert dataset["project_id"] and dataset["current_version_id"]
        totals = client.post(f"/api/v1/datasets/{dataset['id']}/sql", json={
            "workspace_id": workspace["id"],
            "sql": 'SELECT SUM("Sales") AS sales, SUM("Profit") AS profit FROM dataset',
            "max_rows": 10,
        }).json()["rows"][0]
        assert totals == {"sales": 53_819_405.0, "profit": 4_096_303.0}
        regions = client.post(f"/api/v1/datasets/{dataset['id']}/sql", json={
            "workspace_id": workspace["id"],
            "sql": 'SELECT "Region", SUM("Sales") AS sales, SUM("Profit") AS profit FROM dataset GROUP BY 1 ORDER BY 1',
            "max_rows": 10,
        }).json()["rows"]
        assert regions == [
            {"Region": "Central", "sales": 13_460_268.0, "profit": 1_023_950.0},
            {"Region": "East", "sales": 13_449_435.0, "profit": 1_024_202.0},
            {"Region": "South", "sales": 13_473_090.0, "profit": 1_023_940.0},
            {"Region": "West", "sales": 13_436_612.0, "profit": 1_024_211.0},
        ]

        projects = client.get("/api/v1/projects", params={"workspace_id": workspace["id"]}).json()
        assert len(projects) == 1 and projects[0]["is_default"] is True
        project_id = projects[0]["id"]
        versions = client.get(f"/api/v1/projects/{project_id}/dataset-versions").json()
        source_version = next(item for item in versions if item["id"] == dataset["current_version_id"])
        assert source_version["content_hash"] == expected_hash
        assert source_version["row_count"] == 51_290
        assert len(source_version["column_schema"]) == 27

        cleaned_response = client.post(f"/api/v1/datasets/{dataset['id']}/cleaning/apply", json={
            "workspace_id": workspace["id"], "name": "Global Superstore cleaned",
            "steps": [{"operation": "trim_text", "column": "Category"}],
        })
        assert cleaned_response.status_code == 201
        cleaned = cleaned_response.json()
        assert cleaned["recipe"]["source_version_id"] == dataset["current_version_id"]
        cleaned_version = client.get(
            f"/api/v1/dataset-versions/{cleaned['dataset']['current_version_id']}"
        ).json()
        assert cleaned_version["parent_version_id"] == dataset["current_version_id"]

        conversation = client.post("/api/v1/conversations", json={
            "workspace_id": workspace["id"], "dataset_id": cleaned["dataset"]["id"], "title": "Global Superstore 回归",
            "clarification_mode": "off",
        }).json()
        planned = client.post("/api/v1/chat/plan", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "dataset_id": cleaned["dataset"]["id"], "analysis_mode": "business",
            "message": "生成 Global Superstore 管理层报告并说明数据质量",
        }).json()
        queued = client.post("/api/v1/chat/execute-async", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "plan_message_id": planned["user_message"]["id"],
        })
        assert queued.status_code == 202
        run_id = queued.json()["id"]
        duplicate = client.post("/api/v1/chat/execute-async", json={
            "workspace_id": workspace["id"], "conversation_id": conversation["id"],
            "plan_message_id": planned["user_message"]["id"],
        })
        assert duplicate.status_code == 202 and duplicate.json()["id"] == run_id
        for _ in range(240):
            runs = client.get(f"/api/v1/projects/{project_id}/runs").json()
            current = next(item for item in runs if item["id"] == run_id)
            if current["status"] not in {"queued", "running"}:
                break
            time.sleep(0.05)
        assert current["status"] == "completed"
        assert current["dataset_version_ids"] == [cleaned_version["id"]]
        assert current["analysis_spec"]["analysis_mode"] == "business"

        evidence = client.get(f"/api/v1/runs/{run_id}/evidence").json()
        assert {item["evidence_key"] for item in evidence} >= {"dataset-shape", "quality-summary"}
        assert all(item["dataset_version_ids"] == [cleaned_version["id"]] for item in evidence)
        artifacts = client.get(f"/api/v1/projects/{project_id}/artifacts").json()
        report_artifact = next(item for item in artifacts if item["run_id"] == run_id and item["kind"] == "report")
        assert report_artifact["version"] == 1
        document = report_artifact['content']
        report_evidence = {item['id']:item for item in document['evidence']}
        assert report_evidence['metric-sales']['value'] == '53,819,405.00'
        assert report_evidence['metric-profit']['value'] == '4,096,303.00'
        assert report_evidence['group-primary']['data'][0]['Region'] == 'South'
        assert document['quality']['passed']
        assert all(item['status']=='verified' for item in document['quality']['claim_checks'])
        assert report_artifact["source_dataset_version_ids"] == [cleaned_version["id"]]
        assert all(item["artifact_id"] == report_artifact["id"] for item in evidence)
        events = client.get(f"/api/v1/analysis-runs/{run_id}/events", params={"workspace_id": workspace["id"]}).json()
        assert events[0]["event_type"] == "run.queued"
        assert events[-1]["event_type"] == "run.completed"
        assert {item["event_type"] for item in events} >= {"step.started", "step.completed", "evidence.created", "artifact.created"}
        steps = [item["data"] for item in events if item["event_type"] == "step.completed"]
        assert all(item["execution"]["status"] == item["result_status"] for item in steps)
        assert all("input_summary" in item["execution"] and "output_summary" in item["execution"] for item in steps)
        assert any(item["evidence"] for item in steps)
